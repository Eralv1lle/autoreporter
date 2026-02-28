from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

import json

from config import config


def make_docx_from_json(data: dict, user_id: int, num: int = 0) -> str:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = data["styles"]["font"]
    style.font.size = data["styles"]["base-size"]

    aligns = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY
    }

    par = None
    for item in data["content"]:
        if not item["add_run"] or par is None:
            par = doc.add_paragraph()
            par.alignment = aligns[item["align"]]
            r = par.add_run(item["text"])
            r.font.size = Pt(item["size"])
            r.bold = item["bold"]
            r.italic = item["italic"]
            r.underline = item["underlined"]
        else:
            r = par.add_run(item["text"])
            r.font.size = Pt(item["size"])
            r.bold = item["bold"]
            r.italic = item["italic"]
            r.underline = item["underlined"]

    filename = f"{data["report_name"]}_{user_id}_{num}.docx"
    doc.save(f"{config.DOCS_DIRECTORY}/{filename}")
    return filename
